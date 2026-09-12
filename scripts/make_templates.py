import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "templates")
os.makedirs(OUT, exist_ok=True)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="B7C3D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def build_college_schedule():
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"

    days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"]
    headers = ["Day", "P1", "P2", "P3", "P4", "P5", "P6", "P7", "Holiday?"]
    ws.append(headers)
    style_header(ws, 1, len(headers))

    example = {
        "Wednesday": {"P3": "PD: C Programming"},
        "Thursday": {"P3": "PD: English Language"},
        "Friday": {"P3": "PD: Aptitude"},
    }

    for d in days:
        row = [d, "", "", "", "", "", "", "", ""]
        for key, label in (example.get(d) or {}).items():
            idx = int(key[1:]) + 1
            row[idx] = label
        ws.append(row)

    for r in range(2, 8):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal=("center" if c > 1 else "left"), vertical="center")

    ws.column_dimensions["A"].width = 14
    for c in range(2, 9):
        ws.column_dimensions[get_column_letter(c)].width = 18
    ws.column_dimensions["I"].width = 11

    wsi = wb.create_sheet("Instructions")
    notes = [
        ["College Schedule Template — how to fill"],
        ["",
         "1.  Day column: Monday..Saturday (or Sunday if your college works that day)."],
        ["",
         "2.  P1..P7 are the periods of each day. Up to 7 periods per day is supported."],
        ["",
         "3.  Leave a period cell EMPTY  = the period is free and your staff can be scheduled there."],
        ["",
         "4.  Type ANY text in a cell to block it, e.g. 'Lunch break', 'Sports', 'Assembly'."],
        ["",
         "5.  The three professional-development hours are pre-filled as examples on Wednesday (C Programming), Thursday (English) and Friday (Aptitude). Move them to the correct periods for your college, edit the text, or delete them."],
        ["",
         "6.  Put 'YES' in the Holiday? column for a day to make the ENTIRE day unavailable."],
        ["",
         "7.  Delete any day row that is not a working day (for example Saturday) — or mark it Holiday."],
        ["",
         "8.  Save as .xlsx and upload it as the College Schedule file."],
    ]
    for r in notes:
        wsi.append(r)
    wsi.column_dimensions["A"].width = 2
    wsi.column_dimensions["B"].width = 110
    for i in range(1, len(notes) + 1):
        wsi.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    wsi.cell(row=1, column=2).font = Font(bold=True, size=13)

    wb.save(os.path.join(OUT, "college_schedule_template.xlsx"))
    print("wrote templates/college_schedule_template.xlsx")


def build_staff_requests():
    wb = Workbook()
    ws = wb.active
    ws.title = "Requests"

    headers = [
        "Staff Name", "Department", "Subject", "Classes per Week",
        "Class Type", "Duration per Class (periods)", "Pattern", "Preferred Day",
    ]
    ws.append(headers)
    style_header(ws, 1, len(headers))

    sample = ["Jane Doe", "Computer Science", "Data Structures", "2", "Normal", "1", "Separate", "Any"]
    ws.append(sample)

    for r in range(2, 200):
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).border = BORDER

    dv_type = DataValidation(type="list", formula1='"Normal,Lab"', allow_blank=True)
    dv_type.error = "Choose Normal or Lab"
    dv_dur = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True)
    dv_dur.error = "Choose 1, 2 or 3"
    dv_pat = DataValidation(type="list", formula1='"Separate,Continuous"', allow_blank=True)
    dv_pat.error = "Choose Separate or Continuous"
    dv_day = DataValidation(type="list", formula1='"Any,Mon,Tue,Wed,Thu,Fri,Sat"', allow_blank=True)
    dv_day.error = "Choose Any or a weekday"

    ws.add_data_validation(dv_type)
    ws.add_data_validation(dv_dur)
    ws.add_data_validation(dv_pat)
    ws.add_data_validation(dv_day)

    dv_type.add("E2:E200")
    dv_dur.add("F2:F200")
    dv_pat.add("G2:G200")
    dv_day.add("H2:H200")

    widths = [20, 18, 24, 16, 13, 24, 16, 15]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wsi = wb.create_sheet("Instructions")
    notes = [
        ["Staff Requests Template — how to fill"],
        ["",
         "1.  One row per subject a staff member teaches. If a staff member teaches two subjects, add two rows."],
        ["",
         "2.  Staff Name: any staff member who appears before is reused, not duplicated."],
        ["", "3.  Classes per Week: how many times this subject meets in the week (1, 2, 3, ...)."],
        ["",
         "4.  Class Type: Normal (theory) or Lab. Lab classes must be 2 or 3 periods long."],
        ["",
         "5.  Duration per Class (periods): how many consecutive periods one class occupies (1, 2 or 3)."],
        ["",
         "6.  Pattern (used when Classes per Week > 1):"],
        ["", "      Continuous: all classes of this subject sit together as one block (e.g. a 2-hour lab)."],
        ["", "      Separate: each class is placed on a different day."],
        ["",
         "7.  Preferred Day: best-effort hint (Any, Mon, Tue, Wed, Thu, Fri, Sat). If it cannot be granted, the subject is placed on another free day."],
        ["",
         "8.  A cell that is booked once can never be booked by anyone else — no double bookings are possible."],
        ["", "9.  Save as .xlsx and upload it as the Staff Requests file."],
    ]
    for r in notes:
        wsi.append(r)
    wsi.column_dimensions["A"].width = 2
    wsi.column_dimensions["B"].width = 110
    for i in range(1, len(notes) + 1):
        wsi.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    wsi.cell(row=1, column=2).font = Font(bold=True, size=13)

    wb.save(os.path.join(OUT, "staff_requests_template.xlsx"))
    print("wrote templates/staff_requests_template.xlsx")


if __name__ == "__main__":
    build_college_schedule()
    build_staff_requests()